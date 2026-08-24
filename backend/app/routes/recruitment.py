import os
import uuid
from datetime import datetime, timedelta

from flask import Blueprint, request, jsonify, current_app
from flask_jwt_extended import jwt_required, get_jwt_identity, get_jwt
from werkzeug.utils import secure_filename

from app.extensions import db
from app.models import (
    JobPosting, Candidate, Interview, InterviewPanelist, User, Department,
)
from app.utils.security import roles_required
from app.utils.permissions import role_or_permission_required
from app.utils.audit import log_action
from app.utils.dedupe import find_duplicate, normalize_email, normalize_phone
from app.utils.image_gen import generate_job_post_image
from app.roles import RECRUITMENT_ROLES, HR_OPERATIONS, HR_RECRUITER, MANAGER, ADMIN_ROLES

bp = Blueprint("recruitment", __name__, url_prefix="/api/recruitment")

VALID_STAGES = ["applied", "screening", "mock_test", "interview", "offer", "hired", "rejected"]


# ============================================================================
# Job postings
# ============================================================================

@bp.get("/jobs")
@jwt_required()
def list_jobs():
    status = request.args.get("status")
    requisition_type = request.args.get("requisition_type")
    vendor_id = request.args.get("vendor_id", type=int)
    query = JobPosting.query
    if status:
        query = query.filter_by(status=status)
    if requisition_type:
        query = query.filter_by(requisition_type=requisition_type)
    if vendor_id:
        query = query.filter_by(vendor_id=vendor_id)
    rows = query.order_by(JobPosting.created_at.desc()).all()
    return jsonify([j.to_dict() for j in rows])


@bp.get("/jobs/<int:job_id>")
@jwt_required()
def get_job(job_id):
    return jsonify(JobPosting.query.get_or_404(job_id).to_dict())


@bp.post("/jobs")
@roles_required(*RECRUITMENT_ROLES)
def create_job():
    payload = request.get_json(silent=True) or {}
    if not payload.get("title"):
        return jsonify({"message": "Job title is required."}), 400

    requisition_type = payload.get("requisition_type", "internal")
    if requisition_type not in ("internal", "client"):
        return jsonify({"message": "requisition_type must be 'internal' or 'client'."}), 400
    vendor_id = payload.get("vendor_id")
    if requisition_type == "client" and not vendor_id:
        return jsonify({"message": "A client requisition must specify which vendor/client it's for."}), 400
    if requisition_type == "internal":
        vendor_id = None  # never allow a stray vendor_id to leak onto an internal requisition

    job = JobPosting(
        title=payload["title"], department_id=payload.get("department_id"),
        employment_type=payload.get("employment_type", "full_time"), openings=payload.get("openings", 1),
        description=payload.get("description", ""), salary_min=payload.get("salary_min"),
        salary_max=payload.get("salary_max"), priority=payload.get("priority", "standard"),
        auto_close_date=(datetime.strptime(payload["auto_close_date"], "%Y-%m-%d").date()
                          if payload.get("auto_close_date") else None),
        screening_questions=payload.get("screening_questions", ""),
        interview_levels=payload.get("interview_levels", ""),
        requisition_type=requisition_type, vendor_id=vendor_id,
        requisition_status="pending_approval", status="open",
        created_by=int(get_jwt_identity()),
    )
    db.session.add(job)
    db.session.commit()
    return jsonify(job.to_dict()), 201


@bp.put("/jobs/<int:job_id>")
@roles_required(*RECRUITMENT_ROLES)
def update_job(job_id):
    job = JobPosting.query.get_or_404(job_id)
    payload = request.get_json(silent=True) or {}
    for field in ["title", "employment_type", "openings", "description", "status", "department_id",
                   "salary_min", "salary_max", "priority", "screening_questions", "interview_levels",
                   "requisition_type", "vendor_id"]:
        if field in payload:
            setattr(job, field, payload[field])
    if job.requisition_type == "internal":
        job.vendor_id = None
    elif job.requisition_type == "client" and not job.vendor_id:
        return jsonify({"message": "A client requisition must specify which vendor/client it's for."}), 400
    if payload.get("auto_close_date"):
        job.auto_close_date = datetime.strptime(payload["auto_close_date"], "%Y-%m-%d").date()
    db.session.commit()
    return jsonify(job.to_dict())


@bp.post("/jobs/<int:job_id>/approve")
@role_or_permission_required(MANAGER, HR_OPERATIONS, permission_key="approve_job_postings")
def approve_job(job_id):
    job = JobPosting.query.get_or_404(job_id)
    job.requisition_status = "open"
    job.status = "open"
    job.approved_by = int(get_jwt_identity())
    job.approved_at = datetime.utcnow()
    db.session.commit()
    return jsonify(job.to_dict())


@bp.post("/jobs/<int:job_id>/reject")
@role_or_permission_required(MANAGER, HR_OPERATIONS, permission_key="approve_job_postings")
def reject_job(job_id):
    job = JobPosting.query.get_or_404(job_id)
    job.requisition_status = "on_hold"
    job.status = "on_hold"
    db.session.commit()
    return jsonify(job.to_dict())


@bp.post("/jobs/<int:job_id>/clone")
@roles_required(*RECRUITMENT_ROLES)
def clone_job(job_id):
    original = JobPosting.query.get_or_404(job_id)
    clone = JobPosting(
        title=f"{original.title} (Copy)", department_id=original.department_id,
        employment_type=original.employment_type, openings=original.openings,
        description=original.description, salary_min=original.salary_min, salary_max=original.salary_max,
        priority=original.priority, screening_questions=original.screening_questions,
        requisition_status="pending_approval", status="open", created_by=int(get_jwt_identity()),
    )
    db.session.add(clone)
    db.session.commit()
    return jsonify(clone.to_dict()), 201


@bp.post("/jobs/<int:job_id>/generate-post-image")
@roles_required(*RECRUITMENT_ROLES)
def generate_post_image(job_id):
    job = JobPosting.query.get_or_404(job_id)
    payload = request.get_json(silent=True) or {}
    layout = payload.get("layout", "classic-blue")
    filename = f"job_{job.id}_{uuid.uuid4().hex[:8]}.png"
    rel_path = os.path.join("job_posts", filename)
    abs_path = os.path.join(current_app.config["UPLOAD_FOLDER"], rel_path)
    generate_job_post_image(abs_path, job_title=job.title,
                             department=job.department.name if job.department else None,
                             location=payload.get("location", "Bengaluru, India"),
                             experience=payload.get("experience"),
                             apply_link=payload.get("apply_link", f"careers/{job.id}"), layout=layout)
    job.post_image_url = f"/static/uploads/{rel_path}"
    db.session.commit()
    return jsonify({"post_image_url": job.post_image_url})


@bp.post("/jobs/<int:job_id>/post-social")
@roles_required(*RECRUITMENT_ROLES)
def post_social(job_id):
    JobPosting.query.get_or_404(job_id)
    results = {}
    if os.environ.get("LINKEDIN_API_TOKEN"):
        results["linkedin"] = {"status": "not_implemented", "message": "Credentials present but API call not implemented in this build."}
    else:
        results["linkedin"] = {"status": "not_configured", "message": "LINKEDIN_API_TOKEN is not set."}
    if os.environ.get("NAUKRI_API_TOKEN"):
        results["naukri"] = {"status": "not_implemented", "message": "Credentials present but API call not implemented in this build."}
    else:
        results["naukri"] = {"status": "not_configured", "message": "NAUKRI_API_TOKEN is not set."}
    return jsonify(results), 200


# ============================================================================
# Candidate sourcing
# ============================================================================

def _run_auto_validation(candidate):
    reasons = []
    if not candidate.email or "@" not in (candidate.email or ""):
        reasons.append("invalid email format")
    if not candidate.phone or len(normalize_phone(candidate.phone) or "") != 10:
        reasons.append("invalid phone format")
    if not candidate.resume_file_path and not candidate.resume_url:
        reasons.append("no resume attached")
    if not os.environ.get("LINKEDIN_API_TOKEN") and not os.environ.get("NAUKRI_API_TOKEN"):
        reasons.append("external LinkedIn/Naukri cross-check unavailable (no API credentials configured)")
    if reasons:
        candidate.validation_status = "needs_review"
        candidate.validation_notes = "; ".join(reasons)
    else:
        candidate.validation_status = "auto_validated"
        candidate.validation_notes = None


@bp.post("/candidates/check-duplicate")
@roles_required(*RECRUITMENT_ROLES)
def check_duplicate():
    payload = request.get_json(silent=True) or {}
    match = find_duplicate(Candidate, payload.get("email"), payload.get("phone"))
    if not match:
        return jsonify({"duplicate": False})
    return jsonify({"duplicate": True, "candidate": match.to_dict()})


@bp.post("/candidates")
@roles_required(*RECRUITMENT_ROLES)
def add_candidate():
    is_multipart = request.content_type and "multipart/form-data" in request.content_type
    payload = request.form if is_multipart else (request.get_json(silent=True) or {})
    required = ["job_id", "full_name"]
    missing = [f for f in required if not payload.get(f)]
    if missing:
        return jsonify({"message": f"Missing fields: {', '.join(missing)}"}), 400

    email = payload.get("email")
    phone = payload.get("phone")
    duplicate_action = payload.get("duplicate_action")
    match = find_duplicate(Candidate, email, phone)
    if match and duplicate_action != "link_new_application":
        return jsonify({"message": "A candidate with this email or phone already exists.",
                         "duplicate": True, "candidate": match.to_dict()}), 409

    candidate = Candidate(
        job_id=payload["job_id"], full_name=payload["full_name"], email=email,
        email_normalized=normalize_email(email), phone=phone, phone_normalized=normalize_phone(phone),
        resume_url=payload.get("resume_url"), source=payload.get("source"), skills=payload.get("skills"),
        location=payload.get("location"),
        open_to_relocation=str(payload.get("open_to_relocation", "")).lower() in ("true", "1", "on"),
        added_by=int(get_jwt_identity()), assigned_to=int(get_jwt_identity()),
        duplicate_of_candidate_id=match.id if match else None,
        referring_employee_id=payload.get("referring_employee_id") or None,
        referral_bonus_status="pending" if payload.get("referring_employee_id") else None,
    )

    if is_multipart and "resume" in request.files:
        file = request.files["resume"]
        if file and file.filename:
            safe_name = secure_filename(f"{uuid.uuid4().hex[:10]}_{file.filename}")
            rel_path = os.path.join("resumes", safe_name)
            abs_path = os.path.join(current_app.config["UPLOAD_FOLDER"], rel_path)
            os.makedirs(os.path.dirname(abs_path), exist_ok=True)
            file.save(abs_path)
            candidate.resume_file_path = f"/static/uploads/{rel_path}"

    _run_auto_validation(candidate)
    db.session.add(candidate)
    db.session.flush()
    log_action(candidate.id, "created", actor_id=int(get_jwt_identity()),
               detail=f"validation={candidate.validation_status}")
    if match:
        log_action(candidate.id, "linked_as_reapplication", actor_id=int(get_jwt_identity()),
                   detail=f"Linked to existing candidate #{match.id}")
    db.session.commit()
    return jsonify(candidate.to_dict()), 201


@bp.get("/candidates/bulk-import/template")
@jwt_required()
def bulk_import_template():
    import io
    import openpyxl
    from flask import send_file

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Candidates"
    headers = ["name", "email", "phone", "job", "source", "skills", "location", "open_to_relocation"]
    ws.append(headers)
    ws.append(["Jane Doe", "jane@example.com", "9876543210", "Backend Engineer", "linkedin",
               "Java, Spring Boot", "Bengaluru", "yes"])
    for col_idx, header in enumerate(headers, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = max(14, len(header) + 4)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(
        buf, as_attachment=True, download_name="candidate_bulk_import_template.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@bp.post("/candidates/bulk-import/preview")
@roles_required(*RECRUITMENT_ROLES)
def bulk_import_preview():
    if "file" not in request.files:
        return jsonify({"message": "No file uploaded."}), 400
    import openpyxl
    file = request.files["file"]
    wb = openpyxl.load_workbook(file, data_only=True)
    ws = wb.active
    headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]

    def col(row, name):
        try:
            return row[headers.index(name)].value
        except (ValueError, IndexError):
            return None

    valid_rows, flagged_rows, duplicate_rows = [], [], []
    for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
        if all(c.value is None for c in row):
            continue
        name = col(row, "name")
        email = col(row, "email")
        phone = str(col(row, "phone") or "")
        job_ref = col(row, "job")
        source = col(row, "source")
        skills = col(row, "skills")
        location = col(row, "location")
        open_to_relocation = str(col(row, "open_to_relocation") or "").strip().lower() in ("yes", "true", "1")
        reasons = []
        if not name:
            reasons.append("missing name")
        if not email or "@" not in str(email):
            reasons.append("invalid email")
        norm_phone = normalize_phone(phone)
        if not norm_phone or len(norm_phone) != 10:
            reasons.append("invalid phone")
        job = None
        if job_ref:
            job = JobPosting.query.filter(
                db.or_(JobPosting.id == job_ref if str(job_ref).isdigit() else False,
                       JobPosting.title.ilike(str(job_ref)))
            ).first()
        if not job:
            reasons.append("job not found / not matched")
        row_data = {"row": i, "name": name, "email": email, "phone": phone,
                    "job_id": job.id if job else None, "job_title": job.title if job else str(job_ref),
                    "source": source, "skills": skills, "location": location, "open_to_relocation": open_to_relocation}
        match = find_duplicate(Candidate, email, phone)
        if match:
            row_data["existing_candidate"] = match.to_dict()
            duplicate_rows.append(row_data)
        elif reasons:
            row_data["reasons"] = reasons
            flagged_rows.append(row_data)
        else:
            valid_rows.append(row_data)

    return jsonify({"valid_rows": valid_rows, "flagged_rows": flagged_rows, "duplicate_rows": duplicate_rows,
                     "summary": {"valid": len(valid_rows), "flagged": len(flagged_rows), "duplicates": len(duplicate_rows)}})


@bp.post("/candidates/bulk-import/confirm")
@roles_required(*RECRUITMENT_ROLES)
def bulk_import_confirm():
    payload = request.get_json(silent=True) or {}
    rows = payload.get("rows", [])
    actor_id = int(get_jwt_identity())
    created = []
    for row in rows:
        match = find_duplicate(Candidate, row.get("email"), row.get("phone"))
        if match and not row.get("link_duplicate"):
            continue
        candidate = Candidate(
            job_id=row["job_id"], full_name=row["name"], email=row.get("email"),
            email_normalized=normalize_email(row.get("email")), phone=row.get("phone"),
            phone_normalized=normalize_phone(row.get("phone")), source=row.get("source"),
            skills=row.get("skills"), location=row.get("location"),
            open_to_relocation=bool(row.get("open_to_relocation")),
            added_by=actor_id, assigned_to=actor_id,
            duplicate_of_candidate_id=match.id if match else None,
        )
        _run_auto_validation(candidate)
        db.session.add(candidate)
        db.session.flush()
        log_action(candidate.id, "created", actor_id=actor_id, detail="Added via bulk Excel import")
        created.append(candidate)
    db.session.commit()
    return jsonify([c.to_dict() for c in created]), 201


# ============================================================================
# Candidate feed, detail, stage, reassignment, notes, skills
# ============================================================================

@bp.get("/candidates/feed")
@roles_required(MANAGER, HR_OPERATIONS)
def newly_added_feed():
    job_id = request.args.get("job_id", type=int)
    recruiter_id = request.args.get("recruiter_id", type=int)
    days = request.args.get("days", default=14, type=int)
    query = Candidate.query.filter(Candidate.created_at >= datetime.utcnow() - timedelta(days=days))
    if job_id:
        query = query.filter_by(job_id=job_id)
    if recruiter_id:
        query = query.filter_by(added_by=recruiter_id)
    rows = query.order_by(Candidate.created_at.desc()).all()
    return jsonify([c.to_dict() for c in rows])


@bp.get("/jobs/<int:job_id>/candidates")
@jwt_required()
def job_candidates(job_id):
    JobPosting.query.get_or_404(job_id)
    rows = Candidate.query.filter_by(job_id=job_id).order_by(Candidate.created_at.desc()).all()
    return jsonify([c.to_dict() for c in rows])


def _build_candidate_query(args):
    """Shared filter/sort logic for the All Candidates list and its CSV
    export — kept in one place so the two can never drift out of sync."""
    q = args.get("q", "").strip()
    stage = args.get("stage")
    job_id = args.get("job_id", type=int)
    source = args.get("source")
    level = args.get("level")
    requisition_type = args.get("requisition_type")
    sort = args.get("sort", "created_desc")

    query = Candidate.query
    if requisition_type:
        query = query.join(JobPosting).filter(JobPosting.requisition_type == requisition_type)
    if q:
        like = f"%{q}%"
        query = query.filter(
            db.or_(Candidate.full_name.ilike(like), Candidate.email.ilike(like), Candidate.phone.ilike(like))
        )
    if stage:
        query = query.filter_by(stage=stage)
    if job_id:
        query = query.filter_by(job_id=job_id)
    if source:
        query = query.filter_by(source=source)
    if level:
        query = query.filter_by(current_level=level)

    sort_map = {
        "created_desc": Candidate.created_at.desc(),
        "created_asc": Candidate.created_at.asc(),
        "name_asc": Candidate.full_name.asc(),
        "name_desc": Candidate.full_name.desc(),
        "stage_asc": Candidate.stage.asc(),
    }
    if sort == "job_asc":
        if not requisition_type:  # avoid a duplicate join if we already joined above
            query = query.join(JobPosting)
        return query.order_by(JobPosting.title.asc()).all()
    return query.order_by(sort_map.get(sort, Candidate.created_at.desc())).all()


@bp.get("/candidates")
@jwt_required()
def list_all_candidates():
    """Flat, searchable, sortable list of every candidate across every job —
    backs the 'All Candidates' page. Query params:
      q                — matches name, email, or phone (case-insensitive substring)
      stage            — exact stage filter
      job_id           — exact job filter
      source           — exact source filter
      level            — exact current_level filter
      requisition_type — "internal" or "client" — keeps AnametaIOT's own
                          hiring pipeline visually/functionally separate from
                          candidates being sourced for a vendor/client
      sort             — one of: created_desc (default), created_asc, name_asc,
                          name_desc, stage_asc, job_asc
    """
    rows = _build_candidate_query(request.args)
    return jsonify([c.to_dict() for c in rows])


@bp.get("/candidates/export")
@jwt_required()
def export_all_candidates():
    """Same filters as /candidates, returned as a downloadable CSV."""
    import csv
    import io
    from flask import Response

    rows = _build_candidate_query(request.args)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([
        "Name", "Email", "Phone", "Job", "Requisition Type", "Vendor", "Source",
        "Skills", "Location", "Open to Relocation", "Stage", "Current Level",
        "Validation Status", "Verification Status", "Added On",
    ])
    for c in rows:
        d = c.to_dict()
        writer.writerow([
            d["full_name"], d["email"], d["phone"], d.get("job_title"),
            c.job.requisition_type if c.job else "", c.job.vendor.name if (c.job and c.job.vendor) else "",
            d["source"], ", ".join(d["skills"]), d.get("location"),
            "Yes" if d.get("open_to_relocation") else "No", d["stage"], d.get("current_level"),
            d["validation_status"], d["verification_status"], d["created_at"],
        ])
    return Response(
        buf.getvalue(), mimetype="text/csv",
        headers={"Content-Disposition": "attachment; filename=candidates_export.csv"},
    )


@bp.get("/jobs/<int:job_id>/candidates/export")
@jwt_required()
def export_job_candidates(job_id):
    """CSV export for one job's candidate pipeline (used from the Candidates tab)."""
    import csv
    import io
    from flask import Response

    JobPosting.query.get_or_404(job_id)
    rows = Candidate.query.filter_by(job_id=job_id).order_by(Candidate.created_at.desc()).all()
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([
        "Name", "Email", "Phone", "Source", "Skills", "Location", "Open to Relocation",
        "Stage", "Current Level", "Validation Status", "Verification Status", "Added On",
    ])
    for c in rows:
        d = c.to_dict()
        writer.writerow([
            d["full_name"], d["email"], d["phone"], d["source"], ", ".join(d["skills"]),
            d.get("location"), "Yes" if d.get("open_to_relocation") else "No",
            d["stage"], d.get("current_level"), d["validation_status"], d["verification_status"], d["created_at"],
        ])
    return Response(
        buf.getvalue(), mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename=job_{job_id}_candidates.csv"},
    )


@bp.get("/candidates/dashboard-stats")
@jwt_required()
def candidates_dashboard_stats():
    """Feeds a dashboard widget: today's added count + recently status-changed
    candidates, for a quick-actions panel."""
    from app.models import CandidateAuditLog

    today_start = datetime.utcnow().replace(hour=0, minute=0, second=0, microsecond=0)
    added_today = Candidate.query.filter(Candidate.created_at >= today_start).count()

    recent_stage_changes = (
        CandidateAuditLog.query.filter_by(action="stage_changed")
        .order_by(CandidateAuditLog.created_at.desc())
        .limit(8)
        .all()
    )
    recent = []
    for log in recent_stage_changes:
        cand = Candidate.query.get(log.candidate_id)
        if not cand:
            continue
        recent.append({
            "candidate_id": cand.id, "candidate_name": cand.full_name,
            "job_title": cand.job.title if cand.job else None,
            "stage": cand.stage, "detail": log.detail,
            "changed_at": log.created_at.isoformat(),
        })

    return jsonify({
        "added_today": added_today,
        "total_active": Candidate.query.filter(Candidate.stage.notin_(["hired", "rejected"])).count(),
        "recent_stage_changes": recent,
    })


@bp.put("/candidates/<int:candidate_id>/level")
@roles_required(*RECRUITMENT_ROLES)
def update_candidate_level(candidate_id):
    candidate = Candidate.query.get_or_404(candidate_id)
    payload = request.get_json(silent=True) or {}
    level = (payload.get("level") or "").strip()

    if candidate.job and candidate.job.interview_levels:
        allowed = [l.strip() for l in candidate.job.interview_levels.split(",") if l.strip()]
        if level and level not in allowed:
            return jsonify({"message": f"level must be one of: {', '.join(allowed)}"}), 400

    old_level = candidate.current_level
    candidate.current_level = level or None
    log_action(candidate.id, "level_changed", actor_id=int(get_jwt_identity()),
               detail=f"{old_level or '—'} -> {level or '—'}")
    db.session.commit()
    return jsonify(candidate.to_dict())


@bp.get("/candidates/<int:candidate_id>")
@jwt_required()
def get_candidate(candidate_id):
    return jsonify(Candidate.query.get_or_404(candidate_id).to_dict(include_detail=True))


@bp.put("/candidates/<int:candidate_id>/stage")
@roles_required(*RECRUITMENT_ROLES)
def update_stage(candidate_id):
    candidate = Candidate.query.get_or_404(candidate_id)
    payload = request.get_json(silent=True) or {}
    stage = payload.get("stage")
    if stage not in VALID_STAGES:
        return jsonify({"message": f"stage must be one of {VALID_STAGES}"}), 400
    old_stage = candidate.stage
    candidate.stage = stage
    if "notes" in payload:
        candidate.notes = payload["notes"]
    log_action(candidate.id, "stage_changed", actor_id=int(get_jwt_identity()), detail=f"{old_stage} -> {stage}")
    db.session.commit()
    return jsonify(candidate.to_dict())


@bp.put("/candidates/<int:candidate_id>/skills")
@roles_required(*RECRUITMENT_ROLES)
def update_skills(candidate_id):
    candidate = Candidate.query.get_or_404(candidate_id)
    payload = request.get_json(silent=True) or {}
    candidate.skills = payload.get("skills", "")
    log_action(candidate.id, "skills_updated", actor_id=int(get_jwt_identity()), detail=candidate.skills)
    db.session.commit()
    return jsonify(candidate.to_dict())


@bp.post("/candidates/<int:candidate_id>/reassign")
@roles_required(*RECRUITMENT_ROLES)
def reassign_candidate(candidate_id):
    candidate = Candidate.query.get_or_404(candidate_id)
    payload = request.get_json(silent=True) or {}
    new_recruiter_id = payload.get("recruiter_id")
    if not new_recruiter_id:
        return jsonify({"message": "recruiter_id is required."}), 400
    old = candidate.assigned_to
    candidate.assigned_to = new_recruiter_id
    log_action(candidate.id, "reassigned", actor_id=int(get_jwt_identity()), detail=f"recruiter #{old} -> #{new_recruiter_id}")
    db.session.commit()
    return jsonify(candidate.to_dict())


@bp.get("/candidates/<int:candidate_id>/notes")
@jwt_required()
def list_notes(candidate_id):
    Candidate.query.get_or_404(candidate_id)
    from app.models import CandidateNote
    rows = CandidateNote.query.filter_by(candidate_id=candidate_id).order_by(CandidateNote.created_at).all()
    return jsonify([n.to_dict() for n in rows])


@bp.post("/candidates/<int:candidate_id>/notes")
@jwt_required()
def add_note(candidate_id):
    Candidate.query.get_or_404(candidate_id)
    payload = request.get_json(silent=True) or {}
    body = (payload.get("body") or "").strip()
    if not body:
        return jsonify({"message": "Note body is required."}), 400
    from app.models import CandidateNote
    note = CandidateNote(candidate_id=candidate_id, author_id=int(get_jwt_identity()), body=body)
    db.session.add(note)
    db.session.commit()
    return jsonify(note.to_dict()), 201


# ============================================================================
# Verification
# ============================================================================

def _can_verify(candidate, actor):
    if actor.role == HR_OPERATIONS:
        return True
    added_by_user = User.query.get(candidate.added_by)
    if not added_by_user:
        return False
    if actor.id == added_by_user.manager_id:
        return True
    if actor.id == candidate.added_by and added_by_user.self_verify_rights:
        return True
    return False


@bp.post("/candidates/<int:candidate_id>/verify")
@jwt_required()
def verify_candidate(candidate_id):
    candidate = Candidate.query.get_or_404(candidate_id)
    actor_id = int(get_jwt_identity())
    actor = User.query.get_or_404(actor_id)
    if not _can_verify(candidate, actor):
        return jsonify({"message": "You do not have permission to verify this candidate."}), 403
    if candidate.next_revalidation_due and candidate.next_revalidation_due > datetime.utcnow():
        return jsonify({"message": f"Already verified recently. Re-validation opens on {candidate.next_revalidation_due.isoformat()}."}), 400
    recruiter = User.query.get(candidate.added_by)
    interval_days = recruiter.revalidation_interval_days if recruiter else 60
    candidate.verification_status = "verified"
    candidate.verified_by = actor_id
    candidate.verified_at = datetime.utcnow()
    candidate.next_revalidation_due = datetime.utcnow() + timedelta(days=interval_days)
    log_action(candidate.id, "verified", actor_id=actor_id,
               detail=f"Next re-validation due {candidate.next_revalidation_due.date().isoformat()}")
    db.session.commit()
    return jsonify(candidate.to_dict())


@bp.post("/recruiters/<int:user_id>/self-verify")
@roles_required(MANAGER)
def set_self_verify(user_id):
    recruiter = User.query.get_or_404(user_id)
    manager_id = int(get_jwt_identity())
    if recruiter.manager_id != manager_id:
        return jsonify({"message": "You can only manage self-verify rights for your own direct reports."}), 403
    payload = request.get_json(silent=True) or {}
    recruiter.self_verify_rights = bool(payload.get("rights"))
    recruiter.self_verify_granted_by = manager_id if recruiter.self_verify_rights else None
    db.session.commit()
    return jsonify(recruiter.to_dict())


@bp.post("/recruiters/<int:user_id>/revalidation-interval")
@roles_required(MANAGER, HR_OPERATIONS)
def set_revalidation_interval(user_id):
    recruiter = User.query.get_or_404(user_id)
    claims = get_jwt()
    if claims.get("role") == MANAGER and recruiter.manager_id != int(get_jwt_identity()):
        return jsonify({"message": "You can only set this for your own direct reports."}), 403
    payload = request.get_json(silent=True) or {}
    days = payload.get("days")
    if not days or int(days) <= 0:
        return jsonify({"message": "days must be a positive integer."}), 400
    recruiter.revalidation_interval_days = int(days)
    db.session.commit()
    return jsonify(recruiter.to_dict())


# ============================================================================
# Interviews
# ============================================================================

def _has_conflict(interviewer_ids, scheduled_at, exclude_interview_id=None):
    window_start = scheduled_at - timedelta(minutes=30)
    window_end = scheduled_at + timedelta(minutes=30)
    query = (db.session.query(InterviewPanelist).join(Interview).filter(
        InterviewPanelist.interviewer_id.in_(interviewer_ids),
        Interview.scheduled_at.between(window_start, window_end),
        Interview.status.in_(["scheduled", "rescheduled"]),
    ))
    if exclude_interview_id:
        query = query.filter(Interview.id != exclude_interview_id)
    conflict = query.first()
    return conflict.interviewer.full_name if conflict else None


@bp.get("/candidates/<int:candidate_id>/interviews")
@jwt_required()
def list_candidate_interviews(candidate_id):
    Candidate.query.get_or_404(candidate_id)
    rows = Interview.query.filter_by(candidate_id=candidate_id).order_by(Interview.scheduled_at.desc()).all()
    return jsonify([i.to_dict() for i in rows])


@bp.post("/candidates/<int:candidate_id>/interviews")
@roles_required(*RECRUITMENT_ROLES)
def schedule_interview(candidate_id):
    Candidate.query.get_or_404(candidate_id)
    payload = request.get_json(silent=True) or {}
    try:
        scheduled_at = datetime.strptime(payload["scheduled_at"], "%Y-%m-%dT%H:%M")
    except (KeyError, ValueError):
        return jsonify({"message": "scheduled_at is required (format YYYY-MM-DDTHH:MM)."}), 400
    interviewer_ids = payload.get("interviewer_ids", [])
    conflict_name = _has_conflict(interviewer_ids, scheduled_at) if interviewer_ids else None
    interview = Interview(candidate_id=candidate_id, scheduled_at=scheduled_at,
                           round_name=payload.get("round_name", "HR Round"), mode=payload.get("mode", "online"),
                           result="pending", scheduled_by=int(get_jwt_identity()))
    db.session.add(interview)
    db.session.flush()
    for iid in interviewer_ids:
        db.session.add(InterviewPanelist(interview_id=interview.id, interviewer_id=iid))
    log_action(candidate_id, "interview_scheduled", actor_id=int(get_jwt_identity()),
               detail=f"{interview.round_name} at {scheduled_at.isoformat()}")
    db.session.commit()
    response = interview.to_dict()
    if conflict_name:
        response["conflict_warning"] = f"{conflict_name} already has another interview near this time."
    return jsonify(response), 201


@bp.put("/interviews/<int:interview_id>")
@roles_required(*RECRUITMENT_ROLES)
def update_interview(interview_id):
    interview = Interview.query.get_or_404(interview_id)
    payload = request.get_json(silent=True) or {}
    for field in ["feedback", "result", "round_name", "mode", "scorecard"]:
        if field in payload:
            setattr(interview, field, payload[field])
    db.session.commit()
    return jsonify(interview.to_dict())


@bp.post("/interviews/<int:interview_id>/reschedule")
@roles_required(*RECRUITMENT_ROLES)
def reschedule_interview(interview_id):
    interview = Interview.query.get_or_404(interview_id)
    payload = request.get_json(silent=True) or {}
    try:
        new_time = datetime.strptime(payload["scheduled_at"], "%Y-%m-%dT%H:%M")
    except (KeyError, ValueError):
        return jsonify({"message": "scheduled_at is required (format YYYY-MM-DDTHH:MM)."}), 400
    if not payload.get("reason"):
        return jsonify({"message": "A reason is required to reschedule."}), 400
    interviewer_ids = [p.interviewer_id for p in interview.panelists]
    conflict_name = _has_conflict(interviewer_ids, new_time, exclude_interview_id=interview.id)
    interview.scheduled_at = new_time
    interview.status = "rescheduled"
    interview.reschedule_reason = payload["reason"]
    db.session.commit()
    response = interview.to_dict()
    if conflict_name:
        response["conflict_warning"] = f"{conflict_name} already has another interview near this time."
    return jsonify(response)


@bp.post("/interviews/<int:interview_id>/cancel")
@roles_required(*RECRUITMENT_ROLES)
def cancel_interview(interview_id):
    interview = Interview.query.get_or_404(interview_id)
    payload = request.get_json(silent=True) or {}
    interview.status = "cancelled"
    interview.reschedule_reason = payload.get("reason", "")
    db.session.commit()
    return jsonify(interview.to_dict())
