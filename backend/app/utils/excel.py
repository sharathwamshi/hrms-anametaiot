"""Small helper for building .xlsx downloads (openpyxl) with consistent branding.

Every cell that comes from user-controlled text is neutralised against spreadsheet
formula injection (a value starting with = + - @ would otherwise be executed by Excel).
"""
import io
import re
from datetime import datetime, date

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from app.utils.company_time import to_company_local

ORANGE = "E05E00"
ORANGE_LIGHT = "FFF3EA"
MAX_CELL_CHARS = 5000
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def to_local(dt):
    return to_company_local(dt)


def fmt_time(dt):
    local = to_local(dt)
    return f"{local.strftime('%I:%M %p').lstrip('0')} {local.tzname()}" if local else ""


def fmt_datetime(dt):
    local = to_local(dt)
    return f"{local.strftime('%d-%b-%Y %I:%M %p').lstrip('0')} {local.tzname()}" if local else ""


def safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (datetime, date)):
        return value
    text = _ILLEGAL.sub("", str(value))[:MAX_CELL_CHARS]
    if text.startswith(_FORMULA_START):
        text = "'" + text
    return text


def build_workbook(sheets, subtitle=None):
    """sheets: list of {"name", "headers": [...], "rows": [[...]], "note": optional str}.
    Returns BytesIO positioned at 0."""
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="F0D5C0")
    for sheet in sheets:
        ws = wb.create_sheet(re.sub(r"[\\/*?:\[\]]", "-", sheet["name"])[:31] or "Sheet")
        headers, rows = sheet["headers"], sheet["rows"]
        top = 1
        if sheet.get("note") or subtitle:
            ws.cell(row=1, column=1, value=safe_cell(sheet.get("note") or subtitle)).font = Font(italic=True, color="8A5A00", size=10)
            top = 3
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=top, column=c, value=safe_cell(h))
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=ORANGE)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        widths = [min(max(len(str(h)) + 2, 10), 40) for h in headers]
        for r, row in enumerate(rows, top + 1):
            for c, v in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=safe_cell(v))
                cell.border = Border(bottom=thin)
                if (r - top) % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=ORANGE_LIGHT)
                if c <= len(widths):
                    widths[c - 1] = min(max(widths[c - 1], len(str(cell.value)) + 2), 48)
        for c, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = ws.cell(row=top + 1, column=1)
        if rows:
            ws.auto_filter.ref = f"A{top}:{get_column_letter(len(headers))}{top + len(rows)}"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def xlsx_response(sheets, filename, subtitle=None):
    buf = build_workbook(sheets, subtitle)
    safe_name = re.sub(r"[^A-Za-z0-9._-]+", "_", filename)
    if not safe_name.lower().endswith(".xlsx"):
        safe_name += ".xlsx"
    resp = send_file(buf, as_attachment=True, download_name=safe_name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp.headers["Cache-Control"] = "private, no-store"
    return resp
